import os
import uuid
import ast
import re
from io import BytesIO
from datetime import date, datetime
from decimal import Decimal
from flask import Flask, flash, redirect, render_template, request, session, url_for, send_file, send_from_directory
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import inspect, text
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

app = Flask(__name__)
url = os.getenv("DATABASE_URL", "sqlite:///finance.db").replace("postgres://", "postgresql://", 1)
app.config.update(SECRET_KEY=os.getenv("SECRET_KEY", "change-this-before-production"), SQLALCHEMY_DATABASE_URI=url, SQLALCHEMY_TRACK_MODIFICATIONS=False, UPLOAD_FOLDER=os.getenv("UPLOAD_FOLDER", os.path.join(app.root_path,"instance","uploads")), MAX_CONTENT_LENGTH=20*1024*1024)
db = SQLAlchemy(app)
TYPES = {"主合同收入":"contract", "VO变更款项收入":"contract", "分包支出":"subcontract", "其他支出":"other"}
STATUSES = ["未开票", "已开票"]
INCOME_STATUSES = ["未收款", "部分收款", "已收款"]
EXPENSE_STATUSES = ["未支付", "部分支付", "已支付"]
ENTRY_STATUSES = ["进行中", "待处理", "完成", "逾期"]

class User(db.Model):
 id=db.Column(db.Integer,primary_key=True); username=db.Column(db.String(64),unique=True,nullable=False); password_hash=db.Column(db.String(256),nullable=False); is_admin=db.Column(db.Boolean,default=False); created_at=db.Column(db.DateTime,default=datetime.utcnow)
 def set_password(self,p): self.password_hash=generate_password_hash(p)
 def check_password(self,p): return check_password_hash(self.password_hash,p)
class Project(db.Model):
 id=db.Column(db.Integer,primary_key=True); year=db.Column(db.Integer,index=True,nullable=False); name=db.Column(db.String(160),nullable=False); client=db.Column(db.String(160),default=""); manager=db.Column(db.String(80),default=""); contract_amount=db.Column(db.Numeric(14,2),default=0); subcontract_amount=db.Column(db.Numeric(14,2),default=0); other_cost=db.Column(db.Numeric(14,2),default=0); invoice_amount=db.Column(db.Numeric(14,2),default=0); invoice_status=db.Column(db.String(30),default="未开票"); notes=db.Column(db.Text,default=""); updated_at=db.Column(db.DateTime,default=datetime.utcnow,onupdate=datetime.utcnow); created_by=db.Column(db.String(64),default="")
 entries=db.relationship("LedgerEntry",backref="project",cascade="all, delete-orphan",lazy="select",order_by="LedgerEntry.payment_date.desc(), LedgerEntry.id.desc()")
 @property
 def totals(self):
  t={k:Decimal("0") for k in ["contract","subcontract","other","invoice","received","paid","kpi"]}
  for e in self.entries:
   t[e.category]+=e.amount; t["invoice"]+=e.invoice_amount or 0
   if e.category=="contract": t["received"]+=e.received_amount or 0
   else: t["paid"]+=e.received_amount or 0
   t["kpi"]+=sum(e.kpi_by_year.values(),Decimal("0"))
  t["receivable"]=t["contract"];t["receivable_balance"]=t["receivable"]-t["received"]
  t["payable"]=t["subcontract"]+t["other"]-t["paid"]
  t["profit"]=t["contract"]-t["subcontract"]-t["other"]; return t
class LedgerEntry(db.Model):
 id=db.Column(db.Integer,primary_key=True); project_id=db.Column(db.Integer,db.ForeignKey("project.id"),index=True,nullable=False); payment_type=db.Column(db.String(40),nullable=False); category=db.Column(db.String(20),nullable=False); amount=db.Column(db.Numeric(14,2),default=0); currency=db.Column(db.String(8),default="CNY"); entry_status=db.Column(db.String(20),default="进行中"); expected_completion_date=db.Column(db.Date,nullable=True); kpi_included=db.Column(db.Boolean,default=False); kpi_year=db.Column(db.Integer,nullable=True); has_stages=db.Column(db.Boolean,default=False); receipt_status=db.Column(db.String(30),default="未收款"); received_amount=db.Column(db.Numeric(14,2),default=0); invoice_amount=db.Column(db.Numeric(14,2),nullable=True); unissued_invoice_amount=db.Column(db.Numeric(14,2),nullable=True); invoice_status=db.Column(db.String(30),default="未开票"); payment_date=db.Column(db.Date,default=date.today); notes=db.Column(db.Text,default="")
 attachments=db.relationship("Attachment",backref="entry",cascade="all, delete-orphan",lazy="select")
 stages=db.relationship("PaymentStage",backref="entry",cascade="all, delete-orphan",lazy="select",order_by="PaymentStage.sequence")
 @property
 def kpi_by_year(self):
  if self.category!="contract": return {}
  if self.has_stages:
   totals={}
   for stage in self.stages:
    if stage.kpi_included and stage.kpi_year: totals[stage.kpi_year]=totals.get(stage.kpi_year,Decimal("0"))+stage.amount
   return totals
  return {self.kpi_year:self.amount} if self.kpi_included and self.kpi_year else {}
class PaymentStage(db.Model):
 id=db.Column(db.Integer,primary_key=True);entry_id=db.Column(db.Integer,db.ForeignKey("ledger_entry.id"),nullable=False,index=True);sequence=db.Column(db.Integer,nullable=False);amount=db.Column(db.Numeric(14,2),nullable=False,default=0);status=db.Column(db.String(20),default="待处理");kpi_included=db.Column(db.Boolean,default=False);kpi_year=db.Column(db.Integer,nullable=True)
class Attachment(db.Model):
 id=db.Column(db.Integer,primary_key=True); entry_id=db.Column(db.Integer,db.ForeignKey("ledger_entry.id"),nullable=False,index=True); original_name=db.Column(db.String(255),nullable=False); stored_name=db.Column(db.String(255),nullable=False,unique=True); uploaded_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

def me(): return db.session.get(User,session["user_id"]) if session.get("user_id") else None
def login_required(f):
 def w(*a,**k):
  if not me(): flash("请先登录。","warning"); return redirect(url_for("login"))
  return f(*a,**k)
 w.__name__=f.__name__; return w
def admin_required(f):
 def w(*a,**k):
  if not me() or not me().is_admin: flash("此操作仅限管理员。","danger"); return redirect(url_for("dashboard"))
  return f(*a,**k)
 w.__name__=f.__name__; return w
def amount(k):
 try:
  raw=request.form.get(k,"0") or "0"
  raw=re.sub(r"(\d+(?:\.\d+)?)%",r"(\1/100)",raw.replace(" ",""))
  def calculate(node):
   if isinstance(node,ast.Constant) and isinstance(node.value,(int,float)): return Decimal(str(node.value))
   if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):
    value=calculate(node.operand);return value if isinstance(node.op,ast.UAdd) else -value
   if isinstance(node,ast.BinOp) and isinstance(node.op,(ast.Add,ast.Sub,ast.Mult,ast.Div)):
    left,right=calculate(node.left),calculate(node.right)
    if isinstance(node.op,ast.Add): return left+right
    if isinstance(node.op,ast.Sub): return left-right
    if isinstance(node.op,ast.Mult): return left*right
    if right==0: raise ValueError
    return left/right
   raise ValueError
  v=calculate(ast.parse(raw,mode="eval").body)
  if v<0: raise ValueError
  return v
 except: raise ValueError("金额须为大于或等于 0 的数字或计算公式，例如 100000*0.3。")
def optional_amount(k):
 if not request.form.get(k,"").strip(): return None
 return amount(k)
def totals(ps):
 r={k:Decimal("0") for k in ["contract","subcontract","other","invoice","profit"]}
 for p in ps:
  for k,v in p.totals.items(): r[k]+=v
 return r
def export_rows(year):
 rows=[]
 for p in Project.query.filter_by(year=year).order_by(Project.name).all():
  for e in p.entries:
   stages="; ".join(f"阶段{stage.sequence}: {stage.amount} ({stage.status})" for stage in e.stages)
   rows.append([p.name,e.notes or "未命名款项",e.payment_type,e.entry_status,e.payment_date.strftime("%Y-%m-%d"),e.expected_completion_date.strftime("%Y-%m-%d") if e.expected_completion_date else "",e.currency,float(e.amount),e.receipt_status,float(e.received_amount or 0),float(e.invoice_amount) if e.invoice_amount is not None else None,float(e.unissued_invoice_amount) if e.unissued_invoice_amount is not None else None,stages])
 return rows
@app.context_processor
def ctx(): return {"current_user":me(),"now_year":datetime.now().year,"payment_types":TYPES,"invoice_statuses":STATUSES,"income_statuses":INCOME_STATUSES,"expense_statuses":EXPENSE_STATUSES,"settlement_statuses":INCOME_STATUSES+EXPENSE_STATUSES,"entry_statuses":ENTRY_STATUSES}
@app.template_filter("currency")
def currency(v): return f"¥{(v or 0):,.2f}"

@app.route("/login",methods=["GET","POST"])
def login():
 if me(): return redirect(url_for("dashboard"))
 if request.method=="POST":
  u=User.query.filter_by(username=request.form.get("username","").strip()).first()
  if u and u.check_password(request.form.get("password","")): session.clear();session["user_id"]=u.id;return redirect(url_for("dashboard"))
  flash("用户名或密码不正确。","danger")
 return render_template("login.html")
@app.route("/logout")
def logout(): session.clear();return redirect(url_for("login"))
@app.route("/")
@login_required
def dashboard():
 ys=[]
 pending=[]
 for y, in db.session.query(Project.year).distinct().order_by(Project.year.desc()):
  ps=Project.query.filter_by(year=y).all(); ys.append({"year":y,"count":len(ps),"totals":totals(ps)})
 for project in Project.query.order_by(Project.year.desc(),Project.name).all():
  count=sum(1 for entry in project.entries if entry.entry_status=="待处理")
  if count: pending.append({"project":project,"count":count})
 return render_template("dashboard.html",years=ys,pending=pending)
@app.route("/years/<int:year>")
@login_required
def year_detail(year):
 ps=Project.query.filter_by(year=year).order_by(Project.updated_at.desc()).all();return render_template("year_detail.html",year=year,projects=ps,totals=totals(ps))
@app.route("/years/<int:year>/export/<format>")
@login_required
def year_export(year,format):
 headers=["项目","款项","类型","状态","日期","预计完成时间","币种","款项金额","款项状态","已收/已支出","在途开票金额","未开票金额","款项阶段"]; rows=export_rows(year)
 if format=="excel":
  wb=Workbook();ws=wb.active;ws.title=f"{year}年度款项";ws.append(headers)
  for r in rows:ws.append(r)
  for c in ws[1]:c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="EF1746")
  ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(max(max(len(str(x.value or "")) for x in col)+2,10),28)
  buf=BytesIO();wb.save(buf);buf.seek(0);return send_file(buf,as_attachment=True,download_name=f"{year}年度款项明细.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
 if format=="pdf":
  buf=BytesIO();pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"));doc=SimpleDocTemplate(buf,pagesize=landscape(A4),leftMargin=20,rightMargin=20,topMargin=20,bottomMargin=20)
  style=getSampleStyleSheet()["Title"];style.fontName="STSong-Light";style.fontSize=16
  data=[headers]+[["-" if v is None or v=="" else str(v) for v in r] for r in rows];table=Table(data,repeatRows=1,colWidths=[48,58,52,42,48,58,34,52,52,52,55,55,86]);table.setStyle(TableStyle([("FONTNAME",(0,0),(-1,-1),"STSong-Light"),("FONTSIZE",(0,0),(-1,-1),6),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#EF1746")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.25,colors.HexColor("#dddddd")),("VALIGN",(0,0),(-1,-1),"MIDDLE")]))
  doc.build([Paragraph(f"{year}年度款项明细",style),Spacer(1,12),table]);buf.seek(0);return send_file(buf,as_attachment=True,download_name=f"{year}年度款项明细.pdf",mimetype="application/pdf")
 return "格式不支持",400
@app.route("/projects/new",methods=["GET","POST"])
@login_required
def project_new():
 if request.method=="POST":
  try:
   p=Project(year=int(request.form["year"]),name=request.form["name"].strip(),client=request.form.get("client",""),manager=request.form.get("manager",""),notes=request.form.get("notes",""),created_by=me().username)
   if not p.name: raise ValueError("项目名称不能为空。")
   db.session.add(p);db.session.commit();return redirect(url_for("project_detail",project_id=p.id))
  except Exception as e: flash(str(e),"danger")
 return render_template("project_form.html",project=None)
@app.route("/projects/<int:project_id>")
@login_required
def project_detail(project_id):
 p=db.get_or_404(Project,project_id); rank={"待处理":0,"逾期":1,"进行中":2,"完成":3}
 return render_template("project_detail.html",project=p,entries=sorted(p.entries,key=lambda e:(rank.get(e.entry_status,9),e.id)))
@app.route("/projects/<int:project_id>/edit",methods=["GET","POST"])
@login_required
def project_edit(project_id):
 p=db.get_or_404(Project,project_id)
 if request.method=="POST":
  try:
   p.year=int(request.form["year"]);p.name=request.form["name"].strip();p.client=request.form.get("client","");p.manager=request.form.get("manager","");p.notes=request.form.get("notes","")
   if not p.name: raise ValueError("项目名称不能为空。")
   db.session.commit();return redirect(url_for("project_detail",project_id=p.id))
  except Exception as e:flash(str(e),"danger")
 return render_template("project_form.html",project=p)
@app.route("/projects/<int:project_id>/entries/new",methods=["POST"])
@login_required
def entry_new(project_id):
 p=db.get_or_404(Project,project_id);return save_entry(None,p)
@app.route("/entries/<int:entry_id>/edit",methods=["POST"])
@login_required
def entry_edit(entry_id): return save_entry(db.get_or_404(LedgerEntry,entry_id),None)
def save_entry(e,p):
 try:
  typ=request.form["payment_type"]
  if typ not in TYPES:raise ValueError("款项类型无效。")
  if not e:e=LedgerEntry(project_id=p.id)
  e.payment_type=typ;e.category=TYPES[typ];e.amount=amount("receivable_amount");e.currency=request.form.get("currency","CNY");e.entry_status=request.form.get("entry_status","进行中");e.kpi_included=request.form.get("kpi_included")=="on";e.kpi_year=int(request.form["kpi_year"]) if request.form.get("kpi_year") else None;e.has_stages=request.form.get("has_stages")=="on";e.receipt_status=request.form["receipt_status"];e.received_amount=amount("received_amount");e.invoice_amount=optional_amount("in_transit_invoice_amount");e.unissued_invoice_amount=optional_amount("unissued_invoice_amount");e.payment_date=datetime.strptime(request.form["payment_date"],"%Y-%m-%d").date();e.notes=request.form.get("notes","")
  allowed=INCOME_STATUSES if e.category=="contract" else EXPENSE_STATUSES
  if e.entry_status not in ENTRY_STATUSES or e.receipt_status not in allowed:raise ValueError("款项状态与款项类型不匹配。")
  if e.category!="contract": e.kpi_included=False;e.kpi_year=None
  elif e.kpi_included and not e.kpi_year: raise ValueError("计入 KPI 时必须填写年份。")
  if e.has_stages:
   count=int(request.form.get("stage_count",0))
   if count<1 or count>10: raise ValueError("阶段数量需在 1 至 10 之间。")
   e.stages.clear()
   for i in range(1,count+1):
    stage_status=request.form.get(f"stage_status_{i}","待处理")
    if stage_status not in ENTRY_STATUSES: raise ValueError("阶段状态无效。")
    stage_kpi_included=request.form.get(f"stage_kpi_included_{i}")=="on"
    stage_kpi_year=int(request.form[f"stage_kpi_year_{i}"]) if request.form.get(f"stage_kpi_year_{i}") else None
    if e.category!="contract": stage_kpi_included=False;stage_kpi_year=None
    elif stage_kpi_included and not stage_kpi_year: raise ValueError("阶段计入 KPI 时必须填写年份。")
    e.stages.append(PaymentStage(sequence=i,amount=amount(f"stage_amount_{i}"),status=stage_status,kpi_included=stage_kpi_included,kpi_year=stage_kpi_year))
  else: e.stages.clear()
  db.session.add(e);db.session.commit();flash("款项已保存。","success")
 except Exception as x:
  db.session.rollback();flash(str(x),"danger")
 return redirect(url_for("project_detail",project_id=(p or e.project).id))
@app.route("/entries/<int:entry_id>/delete",methods=["POST"])
@login_required
def entry_delete(entry_id):
 e=db.get_or_404(LedgerEntry,entry_id);p=e.project_id
 if not me().check_password(request.form.get("password","")):
  flash("密码验证失败，款项未删除。","danger");return redirect(url_for("project_detail",project_id=p))
 db.session.delete(e);db.session.commit();flash("款项已删除。","success");return redirect(url_for("project_detail",project_id=p))
@app.route("/entries/<int:entry_id>/attachments",methods=["POST"])
@login_required
def attachment_upload(entry_id):
 e=db.get_or_404(LedgerEntry,entry_id);files=[f for f in request.files.getlist("attachment") if f and f.filename]; allowed={"jpg","jpeg","png","gif","webp","pdf","xlsx","xls","doc","docx","csv","txt"}
 if not files: flash("请选择要上传的附件。","danger")
 elif len(files)>5: flash("一次最多上传 5 个附件。","danger")
 elif any("." not in f.filename or f.filename.rsplit(".",1)[1].lower() not in allowed for f in files): flash("包含不支持的文件格式。","danger")
 else:
  os.makedirs(app.config["UPLOAD_FOLDER"],exist_ok=True)
  for f in files:
   original=secure_filename(f.filename) or "attachment"; stored=f"{uuid.uuid4().hex}_{original}";f.save(os.path.join(app.config["UPLOAD_FOLDER"],stored));db.session.add(Attachment(entry=e,original_name=original,stored_name=stored))
  db.session.commit();flash(f"已上传 {len(files)} 个附件。","success")
 return redirect(url_for("project_detail",project_id=e.project_id))
@app.route("/attachments/<int:attachment_id>/download")
@login_required
def attachment_download(attachment_id):
 a=db.get_or_404(Attachment,attachment_id);return send_from_directory(app.config["UPLOAD_FOLDER"],a.stored_name,as_attachment=True,download_name=a.original_name)
@app.route("/attachments/<int:attachment_id>/delete",methods=["POST"])
@login_required
def attachment_delete(attachment_id):
 a=db.get_or_404(Attachment,attachment_id);project_id=a.entry.project_id;path=os.path.join(app.config["UPLOAD_FOLDER"],a.stored_name)
 if os.path.isfile(path): os.remove(path)
 db.session.delete(a);db.session.commit();flash("附件已删除。","success")
 return redirect(url_for("project_detail",project_id=project_id))

@app.route("/users",methods=["GET","POST"])
def users():
 if not me(): return redirect(url_for("login"))
 if request.method=="POST" and me().is_admin:
  name,pw=request.form.get("username",""),request.form.get("password","")
  if len(name)<2 or len(pw)<8: flash("用户名至少 2 位，密码至少 8 位。","danger")
  elif User.query.filter_by(username=name).first(): flash("该用户名已经存在。","danger")
  else:
   u=User(username=name,is_admin=request.form.get("is_admin")=="on");u.set_password(pw);db.session.add(u);db.session.commit();return redirect(url_for("users"))
 return render_template("users.html",users=User.query.order_by(User.created_at.desc()).all() if me().is_admin else [me()])

@app.route("/users/<int:user_id>/update",methods=["POST"])
@login_required
def user_update(user_id):
 target=db.get_or_404(User,user_id)
 if not me().is_admin and target.id != me().id:
  flash("你只能修改自己的账号。","danger");return redirect(url_for("users"))
 name=request.form.get("username","").strip(); password=request.form.get("password","")
 if len(name)<2: flash("用户名至少 2 位。","danger")
 elif User.query.filter(User.username==name,User.id!=target.id).first(): flash("该用户名已经存在。","danger")
 elif password and len(password)<8: flash("新密码至少 8 位。","danger")
 else:
  target.username=name
  if password: target.set_password(password)
  db.session.commit(); flash("账号信息已更新。","success")
 return redirect(url_for("users"))

def initialize():
 db.create_all()
 os.makedirs(app.config["UPLOAD_FOLDER"],exist_ok=True)
 # Add columns for deployments originally created before detailed receipt tracking.
 columns={c['name'] for c in inspect(db.engine).get_columns('ledger_entry')}
 with db.engine.begin() as conn:
  if 'currency' not in columns: conn.execute(text("ALTER TABLE ledger_entry ADD COLUMN currency VARCHAR(8) DEFAULT 'CNY'"))
  if 'receipt_status' not in columns: conn.execute(text("ALTER TABLE ledger_entry ADD COLUMN receipt_status VARCHAR(30) DEFAULT '未收款'"))
  if 'received_amount' not in columns: conn.execute(text("ALTER TABLE ledger_entry ADD COLUMN received_amount NUMERIC(14,2) DEFAULT 0"))
  if 'entry_status' not in columns: conn.execute(text("ALTER TABLE ledger_entry ADD COLUMN entry_status VARCHAR(20) DEFAULT 'ongoing'"))
  if 'has_stages' not in columns: conn.execute(text("ALTER TABLE ledger_entry ADD COLUMN has_stages BOOLEAN DEFAULT FALSE"))
  if 'expected_completion_date' not in columns: conn.execute(text("ALTER TABLE ledger_entry ADD COLUMN expected_completion_date DATE"))
  if 'unissued_invoice_amount' not in columns: conn.execute(text("ALTER TABLE ledger_entry ADD COLUMN unissued_invoice_amount NUMERIC(14,2)"))
  if 'kpi_included' not in columns: conn.execute(text("ALTER TABLE ledger_entry ADD COLUMN kpi_included BOOLEAN DEFAULT FALSE"))
  if 'kpi_year' not in columns: conn.execute(text("ALTER TABLE ledger_entry ADD COLUMN kpi_year INTEGER"))
 stage_columns={c['name'] for c in inspect(db.engine).get_columns('payment_stage')}
 with db.engine.begin() as conn:
  if 'status' not in stage_columns: conn.execute(text("ALTER TABLE payment_stage ADD COLUMN status VARCHAR(20) DEFAULT '待处理'"))
  if 'kpi_included' not in stage_columns: conn.execute(text("ALTER TABLE payment_stage ADD COLUMN kpi_included BOOLEAN DEFAULT FALSE"))
  if 'kpi_year' not in stage_columns: conn.execute(text("ALTER TABLE payment_stage ADD COLUMN kpi_year INTEGER"))
 if not User.query.first():
  u=User(username=os.getenv("ADMIN_USERNAME","admin"),is_admin=True);u.set_password(os.getenv("ADMIN_PASSWORD","ChangeMe123!"));db.session.add(u);db.session.commit()
 for p in Project.query.all():
  if not p.entries and any([p.contract_amount,p.subcontract_amount,p.other_cost]):
   for typ,val,inv,status in [("主合同收入",p.contract_amount,p.invoice_amount,p.invoice_status),("分包支出",p.subcontract_amount,0,"未开票"),("其他支出",p.other_cost,0,"未开票")]:
    if val:db.session.add(LedgerEntry(project=p,payment_type=typ,category=TYPES[typ],amount=val,invoice_amount=inv,invoice_status=status,payment_date=date.today(),notes="由旧版汇总数据转换"))
   p.contract_amount=p.subcontract_amount=p.other_cost=p.invoice_amount=0
 for e in LedgerEntry.query.all():
  if e.entry_status=="ongoing": e.entry_status="进行中"
  if e.entry_status=="Closed": e.entry_status="完成"
  if e.invoice_status not in STATUSES: e.invoice_status="已开票" if e.amount and e.invoice_amount>=e.amount else "未开票"
 db.session.commit()
with app.app_context():initialize()
if __name__=="__main__":app.run(host="0.0.0.0",port=int(os.getenv("PORT",5000)),debug=True)
